"""Excel I/O for batch personalization.

Reads an uploaded .xlsx of prospects and writes back a copy with the
generated personalized email + follow-up appended as new columns. The
original sheet is preserved verbatim — we append columns to the right
so the user's existing layout is never disturbed.

Column detection is fuzzy: we look for case/space-insensitive header
matches against a small synonym table. If a required header is missing
we raise; if optional ones (linkedin) are missing we just skip them.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.levels.schemas import ProspectInput


# Header synonyms — case-insensitive, whitespace-stripped.
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "name":     ("name", "full name", "contact name", "first name + last name", "prospect", "person"),
    "title":    ("title", "job title", "role", "position", "persona"),
    "domain":   ("domain", "company domain", "website", "company website", "url", "company url"),
    "linkedin": ("linkedin", "linkedin url", "linkedin profile", "li", "profile url"),
    "email":    ("email", "email address", "work email"),
}

REQUIRED_HEADERS = ("name", "title", "domain")

# Columns that always come at the tail of the appended block (after the
# per-step Subject/Body pairs).
TRAILING_COLUMNS: tuple[str, ...] = (
    "LinkedIn DM",
    "Model",
    "Slot",
    "Deliverability",
    "Reply Likelihood",
    "Warnings",
)


def _step_columns(max_steps: int) -> list[str]:
    """Build the dynamic per-step header labels.

    Step 1 is the initial email (always there). Steps 2..N are follow-ups
    that exist only when sequence_length>=2 on the job. We don't render
    columns for steps that didn't generate anything on ANY row.
    """
    cols: list[str] = []
    for step in range(1, max_steps + 1):
        cols.append("Subject 1" if step == 1 else f"Subject {step}")
        cols.append("Body 1" if step == 1 else f"Body {step}")
    return cols


def _normalize(s: str | None) -> str:
    return (s or "").strip().lower()


def _build_header_map(headers: list[str]) -> dict[str, int]:
    """Return {logical_field: column_index (1-based)} for headers we can map."""
    mapping: dict[str, int] = {}
    norm = [_normalize(h) for h in headers]
    for field, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:
            if syn in norm:
                mapping[field] = norm.index(syn) + 1  # 1-based for openpyxl
                break
    return mapping


def parse_excel(file_bytes: bytes, sheet_name: str | None = None) -> list[ProspectInput]:
    """Read prospects from the first (or named) sheet of an .xlsx file.

    Raises ValueError if required headers (name, title, domain) are missing.
    """
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    if ws is None or ws.max_row < 2:
        raise ValueError("Sheet is empty or missing data rows")

    headers_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h) if h is not None else "" for h in headers_row]
    col_map = _build_header_map(headers)

    missing = [h for h in REQUIRED_HEADERS if h not in col_map]
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Found headers: {[h for h in headers if h]}. "
            f"Accepted synonyms: " + "; ".join(f"{k}={list(v)}" for k, v in HEADER_SYNONYMS.items() if k in missing)
        )

    prospects: list[ProspectInput] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(field: str) -> str | None:
            idx = col_map.get(field)
            if idx is None or idx > len(row):
                return None
            v = row[idx - 1]
            return None if v is None else str(v).strip() or None

        name = cell("name")
        title = cell("title")
        domain = cell("domain")
        if not (name and title and domain):
            continue  # skip incomplete rows silently
        linkedin = cell("linkedin")
        email = cell("email")
        prospects.append(ProspectInput(
            name=name,
            title=title,
            domain=domain,
            linkedin=linkedin if linkedin else None,
            email=email if email and "@" in email else None,
        ))
    return prospects


def write_excel_with_results(
    original_bytes: bytes,
    results_by_row: dict[int, dict[str, Any]],
    sheet_name: str | None = None,
) -> bytes:
    """Take the original uploaded workbook and append result columns.

    `results_by_row` is keyed by the original Excel row number (2-based,
    matching openpyxl iter_rows row index where the header is row 1).
    Each value is a dict shaped like:

        {
          "steps": [
            {"subject": "...", "body": "..."},   # step 1 (initial)
            {"subject": "...", "body": "..."},   # step 2 (follow-up)
            ...
          ],
          "linkedin": "...",
          "model": "...",
          "slot": "...",
          "deliverability": 87,
          "reply_likelihood": 72,
          "warnings": [...]
        }

    Column count is dynamic — we emit Subject/Body pairs for as many
    steps as the longest sequence actually produced.

    Back-compat: if the old `subject`/`body`/`followup_subject`/
    `followup_body` keys are present (and `steps` is missing), they're
    auto-promoted into a 1-2 element `steps` list before rendering.
    """
    wb = load_workbook(io.BytesIO(original_bytes), data_only=False)
    ws = wb[sheet_name] if sheet_name else wb.active
    if ws is None:
        raise ValueError("No sheet to write to")

    # Normalize back-compat shape → `steps` list everywhere.
    for res in results_by_row.values():
        if "steps" in res:
            continue
        steps: list[dict[str, str]] = []
        if res.get("subject") or res.get("body"):
            steps.append({"subject": res.get("subject", ""), "body": res.get("body", "")})
        if res.get("followup_subject") or res.get("followup_body"):
            steps.append({"subject": res.get("followup_subject", ""), "body": res.get("followup_body", "")})
        res["steps"] = steps

    max_steps = max(
        (len(r.get("steps") or []) for r in results_by_row.values()),
        default=1,
    )
    if max_steps == 0:
        max_steps = 1

    step_cols = _step_columns(max_steps)
    appended_columns = (*step_cols, *TRAILING_COLUMNS)
    first_new_col = ws.max_column + 1

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="6D08BE", end_color="6D08BE", fill_type="solid")  # LakeB2B purple
    header_align = Alignment(horizontal="left", vertical="center")

    for offset, col_label in enumerate(appended_columns):
        col = first_new_col + offset
        c = ws.cell(row=1, column=col, value=col_label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        ws.column_dimensions[get_column_letter(col)].width = (
            32 if "Body" in col_label or col_label == "LinkedIn DM" else 22
        )

    for row_idx, res in results_by_row.items():
        if row_idx < 2:
            continue
        steps = res.get("steps") or []
        values: list[Any] = []
        for step_idx in range(max_steps):
            step = steps[step_idx] if step_idx < len(steps) else {}
            values.extend([step.get("subject", ""), step.get("body", "")])
        values.extend([
            res.get("linkedin", ""),
            res.get("model", ""),
            res.get("slot", ""),
            res.get("deliverability", ""),
            res.get("reply_likelihood", ""),
            ", ".join(res.get("warnings") or []),
        ])
        for offset, value in enumerate(values):
            cell = ws.cell(row=row_idx, column=first_new_col + offset, value=value)
            col_name = appended_columns[offset]
            if "Body" in col_name or col_name == "LinkedIn DM":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
