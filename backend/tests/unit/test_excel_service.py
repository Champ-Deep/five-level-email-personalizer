"""Unit tests for excel_service — focused on the sequence-aware writer.

These assert that the column count is dynamic (one Subject/Body pair per
step actually generated), that the back-compat shape (subject/body +
followup_subject/followup_body) still renders, and that the trailing
fixed columns (LinkedIn DM, Model, Slot, Deliverability, Reply
Likelihood, Warnings) line up correctly after the dynamic block.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel_service import (
    TRAILING_COLUMNS,
    _step_columns,
    write_excel_with_results,
)


def _starter_workbook() -> bytes:
    """A minimal source workbook the writer can append to: 1 header row,
    2 data rows. Matches the shape a user uploads."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws.append(["name", "title", "domain"])
    ws.append(["Priya Sharma", "VP Sales", "stripe.com"])
    ws.append(["Sam Patel", "Founder", "acme.io"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_step_columns_labels():
    assert _step_columns(1) == ["Subject 1", "Body 1"]
    assert _step_columns(3) == [
        "Subject 1", "Body 1",
        "Subject 2", "Body 2",
        "Subject 3", "Body 3",
    ]


def test_write_excel_dynamic_columns_match_max_sequence():
    """Row 1 has a 3-step sequence, row 2 has a 1-step sequence. The
    writer must emit columns for step 3 even though only one row uses
    them — and row 2's step 2/3 cells must be blank, not missing."""
    src = _starter_workbook()
    results = {
        2: {
            "steps": [
                {"subject": "Initial A", "body": "Body A1"},
                {"subject": "Bump A",    "body": "Body A2"},
                {"subject": "Value A",   "body": "Body A3"},
            ],
            "linkedin": "DM A",
            "model": "deepseek/deepseek-v4-pro",
            "slot": "A",
            "deliverability": 88,
            "reply_likelihood": 71,
            "warnings": [],
        },
        3: {
            "steps": [
                {"subject": "Initial B", "body": "Body B1"},
            ],
            "linkedin": "",
            "model": "meta-llama/llama-4-maverick",
            "slot": "B",
            "deliverability": 80,
            "reply_likelihood": 65,
            "warnings": ["body is 64 words (min 65)"],
        },
    }
    out_bytes = write_excel_with_results(src, results)
    wb = load_workbook(io.BytesIO(out_bytes))
    ws = wb.active

    # Headers: name/title/domain + Subject1/Body1/Subject2/Body2/Subject3/Body3 + 6 trailing
    expected_headers = [
        "name", "title", "domain",
        "Subject 1", "Body 1",
        "Subject 2", "Body 2",
        "Subject 3", "Body 3",
        *TRAILING_COLUMNS,
    ]
    actual_headers = [c.value for c in ws[1]]
    assert actual_headers == expected_headers, (
        f"Header mismatch.\n got: {actual_headers}\nwant: {expected_headers}"
    )

    # Row 2 (1st prospect): all 3 steps populated.
    r2 = {c.value for c in ws[2]}
    for v in ["Initial A", "Body A1", "Bump A", "Body A2", "Value A", "Body A3", "DM A", "A", 88, 71]:
        assert v in r2, f"row 2 missing {v!r}; got {r2}"

    # Row 3 (2nd prospect): only step 1 populated; step 2/3 cells blank,
    # warnings cell contains the validation note.
    headers = actual_headers
    row3 = {headers[i]: c.value for i, c in enumerate(ws[3])}
    assert row3["Subject 1"] == "Initial B"
    assert row3["Body 1"] == "Body B1"
    assert row3["Subject 2"] in (None, "")
    assert row3["Body 2"] in (None, "")
    assert row3["Subject 3"] in (None, "")
    assert row3["Body 3"] in (None, "")
    assert row3["Model"] == "meta-llama/llama-4-maverick"
    assert row3["Slot"] == "B"
    assert "min 65" in (row3["Warnings"] or "")


def test_write_excel_back_compat_followup_keys():
    """Old callers (or in-flight Redis state from before the sequence
    migration) pass subject/body + followup_subject/followup_body
    instead of `steps`. The writer must auto-promote that into a
    2-element steps list, yielding Subject1/Body1/Subject2/Body2."""
    src = _starter_workbook()
    results = {
        2: {
            "subject": "Initial via old key",
            "body": "Old-shape body 1",
            "followup_subject": "Old follow-up",
            "followup_body": "Old-shape body 2",
            "linkedin": "",
            "model": "m",
            "slot": "A",
            "deliverability": 80,
            "reply_likelihood": 60,
            "warnings": [],
        },
    }
    out_bytes = write_excel_with_results(src, results)
    wb = load_workbook(io.BytesIO(out_bytes))
    ws = wb.active

    actual_headers = [c.value for c in ws[1]]
    # Should produce 2 step-pairs (initial + follow-up), no step 3.
    assert "Subject 1" in actual_headers
    assert "Body 1" in actual_headers
    assert "Subject 2" in actual_headers
    assert "Body 2" in actual_headers
    assert "Subject 3" not in actual_headers

    row = {actual_headers[i]: c.value for i, c in enumerate(ws[2])}
    assert row["Subject 1"] == "Initial via old key"
    assert row["Body 1"] == "Old-shape body 1"
    assert row["Subject 2"] == "Old follow-up"
    assert row["Body 2"] == "Old-shape body 2"


def test_write_excel_preserves_original_columns():
    """The user's existing columns must come through untouched — we
    only ever append to the right."""
    src = _starter_workbook()
    results = {
        2: {"steps": [{"subject": "S", "body": "B"}], "warnings": []},
        3: {"steps": [{"subject": "S", "body": "B"}], "warnings": []},
    }
    out_bytes = write_excel_with_results(src, results)
    wb = load_workbook(io.BytesIO(out_bytes))
    ws = wb.active

    # First 3 cells of row 1 (header row) are still name/title/domain.
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["name", "title", "domain"]
    # And the data rows still have the original prospect data in those cells.
    assert ws.cell(row=2, column=1).value == "Priya Sharma"
    assert ws.cell(row=3, column=2).value == "Founder"
